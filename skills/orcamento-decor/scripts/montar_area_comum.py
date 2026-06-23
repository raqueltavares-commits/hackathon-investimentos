"""Monta o orçamento de ÁREA COMUM (Recepção, Rooftop, Sinalização, Serviços)
+ Resumo Área Comum (SUBTOTAL + Taxa ADM 13% + Taxa Decor 6% + Budget alvo).

As linhas de cada ambiente vêm do leitor Revit (`revit_para_linhas.py`) + cor dos
renders. Este módulo agrupa por aba, calcula o resumo e serializa pro xlsx.

Diferenças vs unidades:
- ADM 13% incide aqui (na área comum), no Resumo — não por aba.
- Sinalização: "Nome Spot (logo iluminado)" R$2.500 por ambiente (fixo).
- Budget: valor de entrada por Spot (--budget), exibido pra comparar com o total.

Uso:
  python montar_area_comum.py --payload tmp/area_comum.json > tmp/area_comum_out.json
"""
import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from modelos import LinhaMemorial, MemorialTipologia, TAXA_ADM_PCT, TAXA_DECOR_PCT
from montar_orcamento import serializar_estruturado, _linha_total, _brl, N_COLS
from precos_fora_catalogo import marcenaria_por_m2

VALOR_LOGO_ILUMINADO = 2500.0
VALOR_NUMERO_PORTA = 150.0       # número de identificação por porta (1 por unidade)
VALOR_QUADRO_MARKETING = 150.0   # quadro do marketing
QUADROS_POR_PAVIMENTO = 4        # cada pavimento recebe 4 quadros


def linha_sinalizacao(ambiente: str, qtd: int = 1,
                      valor: float = VALOR_LOGO_ILUMINADO) -> LinhaMemorial:
    return LinhaMemorial(
        categoria="SINALIZAÇÃO", ambiente=ambiente,
        item=f"Nome Spot (logo iluminado) — {ambiente}",
        quantidade=qtd, valor_unitario=valor, cor_acabamento="A definir (ver render)",
    )


def linha_numero_porta(qtd_unidades: int, valor: float = VALOR_NUMERO_PORTA) -> LinhaMemorial:
    """1 número de identificação por unidade (qtd = nº de unidades, do Nekt)."""
    return LinhaMemorial(
        categoria="SINALIZAÇÃO", ambiente="Spot",
        item="Números de identificação (porta)", quantidade=int(qtd_unidades),
        valor_unitario=valor, cor_acabamento="A definir",
        referencia="1 por unidade (Nekt total_units)",
    )


def linha_quadros_marketing(qtd_pavimentos: int, valor: float = VALOR_QUADRO_MARKETING) -> LinhaMemorial:
    """4 quadros do marketing por pavimento (qtd = pavimentos × 4)."""
    return LinhaMemorial(
        categoria="SINALIZAÇÃO", ambiente="Spot",
        item="Quadros do marketing", quantidade=int(qtd_pavimentos) * QUADROS_POR_PAVIMENTO,
        valor_unitario=valor, cor_acabamento="—",
        referencia=f"{QUADROS_POR_PAVIMENTO} por pavimento × {int(qtd_pavimentos)} pav.",
    )


def _linha_de_dict(d: dict, ambiente: str = "") -> LinhaMemorial:
    # Marcenaria sem catálogo: área (m²) × R$2.000/m² (qtd=1). Ex: {"marcenaria_m2": 9.76}.
    if d.get("marcenaria_m2"):
        area = float(d["marcenaria_m2"])
        valor = marcenaria_por_m2(area)
        qtd, ref = 1, d.get("referencia") or f"est. {area:g} m² × R$2.000/m²"
    else:
        valor = float(d.get("valor_unitario", 0.0) or 0.0)
        qtd, ref = int(d.get("quantidade", 1) or 1), d.get("referencia", "")
    return LinhaMemorial(
        categoria=d.get("categoria", ""), item=d.get("item", ""),
        quantidade=qtd, valor_unitario=valor,
        referencia=ref,
        ambiente=d.get("ambiente", ambiente),
        fornecedor=d.get("fornecedor", "Catálogo Decor"),
        cor_acabamento=d.get("cor_acabamento", "A definir"),
        frete=float(d.get("frete", 0.0) or 0.0),
    )


def _subtotal(linhas: list[LinhaMemorial]) -> float:
    return round(sum(l.total_linha for l in linhas), 2)


def resumo_area_comum(abas: list[tuple], tem_decor: bool, budget: float | None) -> dict:
    """abas = [(nome, subtotal)]. Aplica ADM 13% + Decor 6% sobre o SUBTOTAL agregado."""
    subtotal = round(sum(s for _n, s in abas), 2)
    taxa_adm = round(subtotal * TAXA_ADM_PCT, 2)
    taxa_decor = round(subtotal * TAXA_DECOR_PCT, 2) if tem_decor else 0.0
    total = round(subtotal + taxa_adm + taxa_decor, 2)
    saldo = round(budget - total, 2) if budget is not None else None
    return {
        "abas": [{"nome": n, "subtotal": s} for n, s in abas],
        "subtotal": subtotal, "taxa_adm": taxa_adm, "taxa_decor": taxa_decor,
        "total": total, "budget": budget, "saldo": saldo,
    }


def _aba_rows(nome: str, spot: str, linhas: list[LinhaMemorial]) -> list[list]:
    """Matriz da aba (itens + TOTAL por categoria, SEM bloco de taxas)."""
    m = MemorialTipologia(tipologia=nome, descricao=nome, estilo="", spot=spot, linhas=linhas)
    return serializar_estruturado(m, "", fechamento=False)


def resumo_rows(spot: str, resumo: dict) -> list[list]:
    """Matriz da aba Resumo Área Comum."""
    rows: list[list] = []
    titulo = [""] * N_COLS
    titulo[0] = f"RESUMO ÁREA COMUM — {spot}"
    rows.append(titulo)
    cab = [""] * N_COLS
    cab[0], cab[N_COLS - 1] = "ABA", "TOTAL (R$)"
    rows.append(cab)
    for a in resumo["abas"]:
        rows.append(_linha_total(a["nome"], a["subtotal"]))
    rows.append([""] * N_COLS)
    rows.append(_linha_total("SUBTOTAL ÁREA COMUM", resumo["subtotal"]))
    rows.append(_linha_total("Taxa Adm (13%)", resumo["taxa_adm"]))
    rows.append(_linha_total("Taxa Decor (6%)", resumo["taxa_decor"]))
    rows.append(_linha_total("TOTAL ÁREA COMUM", resumo["total"]))
    if resumo["budget"] is not None:
        rows.append([""] * N_COLS)
        rows.append(_linha_total("Budget (alvo)", resumo["budget"]))
        rows.append(_linha_total("Saldo vs budget", resumo["saldo"]))
    return rows


def montar_area_comum(payload: dict) -> dict:
    spot = payload.get("spot", "")
    tem_decor = bool(payload.get("tem_decor", True))
    budget = payload.get("budget")
    budget = float(budget) if budget is not None else None

    abas: list[dict] = []
    abas_subtotais: list[tuple] = []

    # 1) Ambientes (Recepção, Rooftop, ...)
    for amb in payload.get("ambientes", []):
        nome = amb["nome"]
        linhas = [_linha_de_dict(d, ambiente=nome) for d in amb.get("linhas", [])]
        abas.append({"nome": nome, "rows": _aba_rows(nome, spot, linhas)})
        abas_subtotais.append((nome, _subtotal(linhas)))

    # 2) Sinalização: logo iluminado por ambiente + números de porta (1/unidade, Nekt)
    #    + quadros do marketing (4/pavimento) + extras informados.
    sinal = payload.get("sinalizacao", [])
    linhas_sin = [linha_sinalizacao(s["ambiente"], int(s.get("qtd", 1) or 1)) for s in sinal]
    if payload.get("total_unidades"):
        linhas_sin.append(linha_numero_porta(int(payload["total_unidades"])))
    if payload.get("total_pavimentos"):
        linhas_sin.append(linha_quadros_marketing(int(payload["total_pavimentos"])))
    linhas_sin += [_linha_de_dict({**d, "categoria": d.get("categoria", "SINALIZAÇÃO")},
                                  ambiente=d.get("ambiente", "Spot"))
                   for d in payload.get("sinalizacao_extras", [])]
    if linhas_sin:
        abas.append({"nome": "Sinalização", "rows": _aba_rows("Sinalização", spot, linhas_sin)})
        abas_subtotais.append(("Sinalização", _subtotal(linhas_sin)))

    # 3) Serviços e Instalações
    servs = payload.get("servicos", [])
    if servs:
        linhas = [_linha_de_dict({**d, "categoria": d.get("categoria", "SERVIÇOS")},
                                 ambiente="Áreas comuns") for d in servs]
        abas.append({"nome": "Serviços e Instalações",
                     "rows": _aba_rows("Serviços e Instalações", spot, linhas)})
        abas_subtotais.append(("Serviços e Instalações", _subtotal(linhas)))

    resumo = resumo_area_comum(abas_subtotais, tem_decor, budget)
    return {"spot": spot, "abas": abas, "resumo": resumo,
            "resumo_rows": resumo_rows(spot, resumo)}


def gerar_xlsx_area_comum(res: dict, spot: str, saida, cabecalho=None) -> list[str]:
    """Workbook de área comum: Resumo Área Comum + uma aba por ambiente/categoria."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from gerar_xlsx import _aba, bold_white, navy_fill

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo Área Comum"
    _aba(ws, res["resumo_rows"], cabecalho=cabecalho)
    ws["A1"].font = Font(bold=True, size=13)

    usados = {"Resumo Área Comum"}
    for a in res["abas"]:
        titulo = a["nome"][:31]
        while titulo in usados:
            titulo = (titulo[:29] + "_2")
        usados.add(titulo)
        ws = wb.create_sheet(title=titulo)
        _aba(ws, a["rows"], cabecalho=cabecalho)
    wb.save(saida)
    return wb.sheetnames


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    p = argparse.ArgumentParser(description="Monta o orçamento de área comum")
    p.add_argument("--payload", type=Path, required=True)
    p.add_argument("--budget", type=float, help="Budget alvo do Spot (sobrescreve o payload)")
    p.add_argument("--saida", type=Path, help="Se informado, gera o .xlsx")
    p.add_argument("--endereco", default="")
    p.add_argument("--cep", default="")
    args = p.parse_args()
    payload = json.loads(args.payload.read_text(encoding="utf-8"))
    if args.budget is not None:
        payload["budget"] = args.budget
    res = montar_area_comum(payload)
    if args.saida:
        cabecalho = [[f"Empreendimento: {res['spot']}"]]
        if args.endereco:
            cabecalho.append([f"Endereço: {args.endereco}"])
        if args.cep:
            cabecalho.append([f"CEP: {args.cep}"])
        abas = gerar_xlsx_area_comum(res, res["spot"], args.saida, cabecalho=cabecalho)
        print(f"xlsx salvo: {args.saida} | abas: {abas}")
    else:
        print(json.dumps(res, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
